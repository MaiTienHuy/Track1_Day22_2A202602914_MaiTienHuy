import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('Day22-AI-Product-GTM-Monetization-Model.xlsx', data_only=False)
print("Sheet names:", wb.sheetnames)

for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n==================== SHEET: {name} (max_row={ws.max_row}, max_col={ws.max_column}) ====================")
    for r in range(1, ws.max_row + 1):
        row_vals = []
        has_content = False
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            val = cell.value
            fill = cell.fill.start_color.rgb if cell.fill and cell.fill.start_color else None
            if val is not None:
                has_content = True
                row_vals.append(f"C{c}[{cell.coordinate}]={val} (fill={fill})")
        if has_content:
            print(f"Row {r:2d}: " + " | ".join(row_vals))

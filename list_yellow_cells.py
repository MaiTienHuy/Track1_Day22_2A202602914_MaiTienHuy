import openpyxl

wb = openpyxl.load_workbook('Day22-AI-Product-GTM-Monetization-Model.xlsx')

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"\n--- {sheetname} ---")
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            # check if yellow
            fill = cell.fill.start_color.rgb if cell.fill and cell.fill.start_color else ""
            if fill == "FFFFF2CC":
                # get label from col A or nearby
                label = ws.cell(r, 1).value or ""
                print(f"YELLOW [{cell.coordinate}] (Row {r}, Col {c}) | Label: {label} | Val: {cell.value}")

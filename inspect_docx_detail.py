import docx

doc = docx.Document('Day22-AI-Product-GTM-One-Pager-Template.docx')

print("=== ALL PARAGRAPHS WITH INDEXES ===")
for i, p in enumerate(doc.paragraphs):
    print(f"P{i:2d} (text len={len(p.text)}): {p.text}")

print("\n=== ALL TABLES ===")
for t_idx, table in enumerate(doc.tables):
    print(f"\n--- Table {t_idx} (rows={len(table.rows)}, cols={len(table.columns)}) ---")
    for r_idx, row in enumerate(table.rows):
        cells_text = [c.text.strip().replace('\n', ' // ') for c in row.cells]
        print(f"  R{r_idx}: " + " | ".join(cells_text))

import docx
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

doc = docx.Document('Day22-AI-Product-GTM-One-Pager-Template.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

print("\n--- PARAGRAPHS ---")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"P{i}: {p.text}")

print("\n--- TABLES ---")
for t_idx, table in enumerate(doc.tables):
    print(f"\nTable {t_idx} (rows={len(table.rows)}, cols={len(table.columns)}):")
    for r_idx, row in enumerate(table.rows):
        row_cells = [cell.text.strip().replace('\n', ' / ') for cell in row.cells]
        print(f"  R{r_idx}: " + " | ".join(row_cells))
